"""文档解析：PDF / DOCX / TXT / MD → 纯文本"""
import re
import uuid
from pathlib import Path

from app.config import settings
from app.utils.text_splitter import split_text

SUPPORTED_TYPES = {"pdf", "docx", "txt", "md"}
MAX_UPLOAD_SIZE = 20 * 1024 * 1024  # 20MB


def parse_file(path: str | Path, file_type: str) -> str:
    """解析文件为纯文本"""
    path = Path(path)
    if file_type == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        pages = []
        for page in reader.pages:
            text = page.extract_text() or ""
            if text.strip():
                pages.append(text)
        return _normalize("\n".join(pages))

    if file_type == "docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                line = " | ".join(c.text.strip() for c in row.cells if c.text.strip())
                if line:
                    parts.append(line)
        return _normalize("\n".join(parts))

    # txt / md
    return _normalize(path.read_text(encoding="utf-8", errors="ignore"))


def save_upload(filename: str, data: bytes) -> Path:
    """保存上传文件到 uploads 目录，返回保存路径"""
    ext = (filename or "").rsplit(".", 1)[-1].lower()
    if ext not in SUPPORTED_TYPES:
        raise ValueError(f"不支持的文件类型，仅支持：{', '.join(sorted(SUPPORTED_TYPES))}")
    if not data:
        raise ValueError("文件内容为空")
    if len(data) > MAX_UPLOAD_SIZE:
        raise ValueError("文件大小不能超过 20MB")
    safe_name = re.sub(r"[^\w\u4e00-\u9fa5.\-]", "_", filename)[:80]
    save_path = settings.uploads_dir / f"{uuid.uuid4().hex[:12]}_{safe_name}"
    save_path.write_bytes(data)
    return save_path


def process_upload(filename: str, data: bytes) -> dict:
    """保存并解析上传文件，返回文本切片。同步实现，请在调用方用 to_thread 包装。"""
    ext = (filename or "").rsplit(".", 1)[-1].lower()
    path = save_upload(filename, data)
    text = parse_file(path, ext)
    chunks = split_text(text, settings.chunk_size, settings.chunk_overlap)
    return {"path": path, "text": text, "chunks": chunks, "file_type": ext}


def _normalize(text: str) -> str:
    """压缩多余空行，规整文本"""
    lines = [line.rstrip() for line in text.splitlines()]
    result: list[str] = []
    blank = 0
    for line in lines:
        if not line.strip():
            blank += 1
            if blank <= 1:
                result.append("")
        else:
            blank = 0
            result.append(line)
    return "\n".join(result).strip()
